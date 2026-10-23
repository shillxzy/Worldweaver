"""DOCX import/export.

Declared support (first version): paragraphs, headings (Heading 1-3 / Title), bold,
italic, underline, strike-through, line breaks, simple bulleted/numbered lists (as
plain list items). Tables, footnotes, comments, headers/footers, images, text boxes and
tracked changes are NOT converted; their text may be skipped. Macros are never
executed: python-docx only reads the XML parts of the package, and ``.docm`` files are
rejected by content type.
"""
import io
import re
import zipfile

from django.conf import settings
from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.exceptions import PackageNotFoundError
from docx.shared import Pt
from rest_framework.exceptions import ValidationError

DOCX_MAIN_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
HEADING_RE = re.compile(r"^(heading|заголовок)\s*(\d)$", re.IGNORECASE)
TITLE_LEVEL = 0


class DocxError(ValidationError):
    pass


def _check_package(raw: bytes):
    if len(raw) > settings.WW_MAX_DOCX_BYTES:
        raise DocxError({"file": f"Файл завеликий (максимум {settings.WW_MAX_DOCX_BYTES // 1024 // 1024} МБ)."})
    if len(raw) == 0:
        raise DocxError({"file": "Файл порожній."})
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            total = sum(info.file_size for info in archive.infolist())
            if total > settings.WW_MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocxError({"file": "Розпакований документ завеликий."})
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names:
                raise DocxError({"file": "Файл не є документом Word (.docx)."})
            content_types = archive.read("[Content_Types].xml").decode("utf-8", "replace")
            if DOCX_MAIN_TYPE not in content_types:
                raise DocxError({"file": "Підтримуються лише документи .docx без макросів."})
    except zipfile.BadZipFile as exc:
        raise DocxError({"file": "Файл пошкоджений або не є .docx."}) from exc


def _heading_level(paragraph):
    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    if style_name.lower() in ("title", "назва"):
        return TITLE_LEVEL
    match = HEADING_RE.match(style_name.strip())
    if match:
        return max(1, min(3, int(match.group(2))))
    return None


def _is_list(paragraph):
    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    if "list" in style_name.lower():
        return True
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def _runs_to_inline(paragraph):
    inline = []
    for run in paragraph.runs:
        marks = []
        if run.bold:
            marks.append({"type": "bold"})
        if run.italic:
            marks.append({"type": "italic"})
        if run.underline:
            marks.append({"type": "underline"})
        if run.font is not None and run.font.strike:
            marks.append({"type": "strike"})
        parts = run.text.replace("\r", "").split("\n")
        for i, part in enumerate(parts):
            if i:
                inline.append({"type": "hardBreak"})
            if part:
                node = {"type": "text", "text": part.replace("\t", "    ")}
                if marks:
                    node["marks"] = [dict(m) for m in marks]
                if inline and inline[-1].get("type") == "text" and inline[-1].get("marks") == node.get("marks"):
                    inline[-1]["text"] += node["text"]
                else:
                    inline.append(node)
    return inline


def parse_docx(raw: bytes):
    """Return a list of blocks: dicts with type 'heading'(level)/'paragraph'/'list_item' and inline content."""
    _check_package(raw)
    try:
        document = Document(io.BytesIO(raw))
    except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise DocxError({"file": "Не вдалося прочитати документ: файл пошкоджений."}) from exc
    except Exception as exc:  # python-docx raises assorted lxml/parsing errors on broken XML
        raise DocxError({"file": "Не вдалося прочитати документ: некоректна структура."}) from exc

    blocks = []
    for paragraph in document.paragraphs:
        inline = _runs_to_inline(paragraph)
        has_text = any(n["type"] == "text" and n["text"].strip() for n in inline)
        if not has_text:
            continue
        level = _heading_level(paragraph)
        if level == TITLE_LEVEL:
            # The document title (Word "Title" style) is shown as a level-1 heading but never splits chapters.
            blocks.append({"type": "heading", "level": 1, "title": True, "content": inline})
        elif level:
            blocks.append({"type": "heading", "level": level, "content": inline})
        elif _is_list(paragraph):
            blocks.append({"type": "list_item", "content": inline})
        else:
            blocks.append({"type": "paragraph", "content": inline})
    if not blocks:
        raise DocxError({"file": "Документ не містить тексту."})
    return blocks


def blocks_to_document(blocks):
    content = []
    pending_list = None
    for block in blocks:
        if block["type"] == "list_item":
            if pending_list is None:
                pending_list = {"type": "bulletList", "content": []}
                content.append(pending_list)
            pending_list["content"].append({"type": "listItem", "content": [{"type": "paragraph", "content": block["content"]}]})
            continue
        pending_list = None
        if block["type"] == "heading":
            content.append({"type": "heading", "attrs": {"level": block["level"]}, "content": block["content"]})
        else:
            content.append({"type": "paragraph", "content": block["content"]})
    return {"type": "doc", "content": content or [{"type": "paragraph"}]}


def _plain(inline):
    return "".join(n.get("text", "") for n in inline).strip()


def _is_split_heading(block, level):
    return block["type"] == "heading" and not block.get("title") and block["level"] == level


def heading_levels(blocks):
    """Heading levels that could split chapters (the Title style is excluded), with counts and samples."""
    stats = {}
    for block in blocks:
        if block["type"] != "heading" or block.get("title"):
            continue
        entry = stats.setdefault(block["level"], {"level": block["level"], "count": 0, "samples": []})
        entry["count"] += 1
        if len(entry["samples"]) < 3:
            entry["samples"].append(_plain(block["content"])[:80])
    return [stats[level] for level in sorted(stats)]


def choose_split_level(blocks):
    """The highest heading level used at least twice; a single heading of a level is usually a book or part title."""
    levels = heading_levels(blocks)
    repeated = [entry["level"] for entry in levels if entry["count"] >= 2]
    if repeated:
        return repeated[0]
    return levels[0]["level"] if levels else None


def plan_chapters(blocks, default_title, split_level=None):
    """Split blocks into chapters on headings of `split_level` (auto-detected when None).

    - Title-style paragraphs never split; a lead-in made only of them is dropped (it is the book title).
    - A split heading followed directly by another one (an arc / part heading with no text of its own)
      does not become an empty chapter: it is kept as an inner heading at the top of the next chapter.
    Returns a list of dicts: {title, blocks, merged: [carried heading titles], lead_in?: front matter}.
    """
    if split_level is None:
        split_level = choose_split_level(blocks)
    if split_level is None or not any(_is_split_heading(b, split_level) for b in blocks):
        return [{"title": default_title, "blocks": list(blocks), "merged": []}]

    chapters, title, current, carried = [], None, [], []

    def flush():
        body = list(current)
        if title is None:
            titles = [_plain(b["content"]) for b in body if b["type"] == "heading" and b.get("title")]
            if len(titles) == len(body):
                return  # only the book title before the first chapter
            # Front matter (epigraph, dedication...) becomes its own opening part named after the book title.
            chapters.append({"title": (titles[0] if titles else "")[:200] or default_title, "blocks": body, "merged": [], "lead_in": True})
            return
        chapters.append({"title": title["text"], "blocks": title["lead"] + body, "merged": title["merged"]})

    for block in blocks:
        if not _is_split_heading(block, split_level):
            current.append(block)
            continue
        text = _plain(block["content"])[:200] or default_title
        if title is not None and not current:
            # The previous heading has no body: carry it into this chapter instead of making an empty one.
            carried.extend(title["lead"])
            carried.append({"type": "heading", "level": split_level, "content": title["heading"]})
            merged = title["merged"] + [title["text"]]
        else:
            flush()
            merged = []
        title = {"text": text, "heading": block["content"], "lead": carried, "merged": merged}
        carried, current = [], []
    flush()
    return chapters


# --- export -------------------------------------------------------------------

def _add_inline(paragraph, nodes):
    for node in nodes or []:
        if node["type"] == "hardBreak":
            if paragraph.runs:
                paragraph.runs[-1].add_break(WD_BREAK.LINE)
            else:
                paragraph.add_run().add_break(WD_BREAK.LINE)
            continue
        if node["type"] != "text":
            continue
        run = paragraph.add_run(node["text"])
        marks = {m["type"] for m in node.get("marks", [])}
        run.bold = "bold" in marks or None
        run.italic = "italic" in marks or None
        run.underline = "underline" in marks or None
        if "strike" in marks:
            run.font.strike = True


def _write_blocks(document, node, list_style=None):
    node_type = node["type"]
    if node_type == "heading":
        level = node.get("attrs", {}).get("level", 2)
        paragraph = document.add_heading(level=min(3, level + 1))
        _add_inline(paragraph, node.get("content"))
    elif node_type == "paragraph":
        paragraph = document.add_paragraph(style=list_style) if list_style else document.add_paragraph()
        _add_inline(paragraph, node.get("content"))
    elif node_type == "horizontalRule":
        paragraph = document.add_paragraph("* * *")
        paragraph.alignment = 1
    elif node_type in ("bulletList", "orderedList"):
        style = "List Bullet" if node_type == "bulletList" else "List Number"
        for item in node.get("content", []):
            for child in item.get("content", []):
                _write_blocks(document, child, style)
    elif node_type == "blockquote":
        for child in node.get("content", []):
            _write_blocks(document, child, "Quote" if child["type"] == "paragraph" else None)
    else:
        for child in node.get("content", []):
            _write_blocks(document, child, list_style)


def export_docx(title, chapters):
    """chapters: iterable of (chapter_title, tiptap_document). Returns bytes."""
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    if title:
        document.add_heading(title, level=0)
    for index, (chapter_title, content) in enumerate(chapters):
        if index and title:
            document.add_page_break()
        document.add_heading(chapter_title, level=1)
        _write_blocks(document, content)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
