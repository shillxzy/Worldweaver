import io

import pytest
from docx import Document

from apps.books.models import Book, Chapter, ChapterVersion

from .conftest import doc

pytestmark = pytest.mark.django_db


def chapters_url(book):
    return f"/api/v1/books/{book.id}/chapters/"


def test_create_book_with_new_world_and_stats(client, user):
    r = client.post("/api/v1/books/", {"title": "Катарсис розуму", "genre": "Фантастика", "new_world_name": "Інститут"}, format="json")
    assert r.status_code == 201, r.content
    data = r.json()
    assert data["world_name"] == "Інститут" and data["progress"] == 0
    assert client.post("/api/v1/books/", {"title": "Без світу"}, format="json").status_code == 400
    book = Book.objects.get(pk=data["id"])
    Chapter.objects.create(book=book, title="1", word_count=100, status="done")
    Chapter.objects.create(book=book, title="2", word_count=50)
    data = client.get(f"/api/v1/books/{book.id}/").json()
    assert (data["chapters_count"], data["word_count"], data["progress"]) == (2, 150, 50)


def test_chapter_crud_reorder_and_word_count(client, book):
    ids = [client.post(chapters_url(book), {"title": f"Розділ {i}"}, format="json").json()["id"] for i in range(3)]
    listing = client.get(chapters_url(book)).json()["results"]
    assert [c["order"] for c in listing] == [0, 1, 2] and "content" not in listing[0]
    r = client.post(chapters_url(book) + "reorder/", {"order": list(reversed(ids))}, format="json")
    assert [c["id"] for c in r.json()] == list(reversed(ids))
    assert client.post(chapters_url(book) + "reorder/", {"order": ids[:2]}, format="json").status_code == 400
    r = client.patch(f"{chapters_url(book)}{ids[0]}/", {"content": doc("Пам’ять — це море.", "Ейра мовчала")}, format="json")
    assert r.status_code == 200, r.content
    assert r.json()["word_count"] == 5 and r.json()["revision"] == 2
    assert "Пам’ять — це море." in r.json()["plain_text"]
    assert client.get(chapters_url(book) + "?search=море").json()["count"] == 1
    assert client.delete(f"{chapters_url(book)}{ids[1]}/").status_code == 204


@pytest.mark.parametrize("bad", [
    {"type": "paragraph"},
    {"type": "doc", "content": [{"type": "script", "text": "x"}]},
    {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "x", "marks": [{"type": "link", "attrs": {"href": "javascript:alert(1)"}}]}]}]},
    {"type": "doc", "content": [{"type": "heading", "attrs": {"level": 9}, "content": []}]},
    {"type": "doc", "content": "not a list"},
    {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": ""}]}]},
])
def test_document_validation_rejects_untrusted_structures(client, chapter, bad):
    r = client.patch(f"{chapters_url(chapter.book)}{chapter.id}/", {"content": bad}, format="json")
    assert r.status_code == 400


def test_document_validation_drops_unknown_attrs_and_limits_size(client, chapter, settings):
    content = {"type": "doc", "content": [{"type": "paragraph", "attrs": {"onclick": "evil()"}, "content": [{"type": "text", "text": "Ok"}]}]}
    r = client.patch(f"{chapters_url(chapter.book)}{chapter.id}/", {"content": content}, format="json")
    assert r.status_code == 200 and "attrs" not in r.json()["content"]["content"][0]
    settings.WW_MAX_DOCUMENT_BYTES = 200
    big = doc("слово " * 100)
    assert client.patch(f"{chapters_url(chapter.book)}{chapter.id}/", {"content": big}, format="json").status_code == 400


def test_revision_conflict_and_versions(client, chapter, settings):
    settings.WW_VERSION_INTERVAL_SECONDS = 0
    url = f"{chapters_url(chapter.book)}{chapter.id}/"
    r = client.patch(url, {"content": doc("Версія два."), "base_revision": 1}, format="json")
    assert r.status_code == 200 and r.json()["revision"] == 2
    r = client.patch(url, {"content": doc("Застаріле збереження."), "base_revision": 1}, format="json")
    assert r.status_code == 409 and r.json()["error"]["code"] == "conflict"
    client.patch(url, {"content": doc("Версія три."), "base_revision": 2, "snapshot": True}, format="json")
    versions = client.get(url + "versions/").json()
    assert len(versions) == 2
    original = next(v for v in versions if v["revision"] == 1)
    assert "Туман" in client.get(url + f"versions/{original['id']}/").json()["plain_text"]
    r = client.post(url + f"versions/{original['id']}/restore/")
    assert r.status_code == 200 and "Туман" in r.json()["plain_text"] and r.json()["revision"] == 4
    assert ChapterVersion.objects.filter(chapter=chapter, reason="restore").exists()


def make_docx(paragraphs):
    document = Document()
    for style, text, bold in paragraphs:
        p = document.add_paragraph(style=style) if style else document.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    buffer.name = "rukopys.docx"
    return buffer


def test_docx_import_split_by_headings_and_formatting(client, book):
    file = make_docx([
        ("Heading 1", "Розділ перший", False), (None, "Туман сповзав.", False), (None, "Жирний текст", True),
        ("List Bullet", "пункт", False), ("Heading 1", "Розділ другий", False), (None, "Світло кликало.", False),
    ])
    r = client.post("/api/v1/import/docx/", {"file": file, "book": book.id, "mode": "split_headings"}, format="multipart")
    assert r.status_code == 201, r.content
    assert [c["title"] for c in r.json()] == ["Розділ перший", "Розділ другий"]
    first = Chapter.objects.get(pk=r.json()[0]["id"])
    para = first.content["content"][1]
    assert para["content"][0]["marks"] == [{"type": "bold"}]
    assert first.content["content"][2]["type"] == "bulletList"


def novel_docx():
    """Shaped like a real manuscript: Title-style book name, Heading 2 for prologue, arcs and chapters."""
    return make_docx([
        ("Title", "СЯЙВО БЕЗОДНІ", False),
        ("Heading 2", "ПРОЛОГ", False), (None, "Темрява дихала.", False),
        ("Heading 2", "Арка: Перший шепіт бурі", False),
        ("Heading 2", "РОЗДІЛ 1", False), (None, "Буря почалась.", False),
        ("Heading 2", "РОЗДІЛ 2", False), (None, "Вітер стих.", False),
        ("Heading 2", "Арка: Азартна гра", False),
        ("Heading 2", "РОЗДІЛ 3", False), (None, "Карти лягли на стіл.", False),
        ("Heading 2", "ЕПІЛОГ", False), (None, "Кінець.", False),
    ])


def test_docx_split_ignores_book_title_and_merges_arcs(client, book):
    r = client.post("/api/v1/import/docx/", {"file": novel_docx(), "book": book.id, "mode": "split_headings"}, format="multipart")
    assert r.status_code == 201, r.content
    assert [c["title"] for c in r.json()] == ["ПРОЛОГ", "РОЗДІЛ 1", "РОЗДІЛ 2", "РОЗДІЛ 3", "ЕПІЛОГ"]
    first_chapter = Chapter.objects.get(pk=r.json()[1]["id"])
    arc, body = first_chapter.content["content"][:2]
    assert arc["type"] == "heading" and arc["content"][0]["text"] == "Арка: Перший шепіт бурі"
    assert body["content"][0]["text"] == "Буря почалась."
    assert "СЯЙВО" not in "".join(c.plain_text for c in Chapter.objects.filter(book=book))


def test_docx_preview_reports_levels_and_chapters(client):
    r = client.post("/api/v1/import/docx/preview/", {"file": novel_docx()}, format="multipart")
    assert r.status_code == 200, r.content
    data = r.json()
    assert data["split_level"] == data["suggested_level"] == 2
    assert [(lv["level"], lv["count"]) for lv in data["levels"]] == [(2, 7)]
    assert [c["title"] for c in data["chapters"]] == ["ПРОЛОГ", "РОЗДІЛ 1", "РОЗДІЛ 2", "РОЗДІЛ 3", "ЕПІЛОГ"]
    assert data["chapters"][1]["merged"] == ["Арка: Перший шепіт бурі"]
    assert data["chapters"][1]["word_count"] > 2  # the carried arc heading counts toward the chapter text


def test_docx_split_level_override_and_single_top_heading(client, book):
    # One Heading 1 (a part title) over several Heading 2 chapters: auto split uses level 2, the lead-in keeps the part title.
    file = make_docx([
        ("Heading 1", "Частина перша", False), (None, "Вступ.", False),
        ("Heading 2", "Глава А", False), (None, "Текст А.", False),
        ("Heading 2", "Глава Б", False), (None, "Текст Б.", False),
    ])
    preview = client.post("/api/v1/import/docx/preview/", {"file": file}, format="multipart").json()
    assert preview["split_level"] == 2
    assert [c["title"] for c in preview["chapters"]] == ["rukopys", "Глава А", "Глава Б"]
    file.seek(0)
    r = client.post("/api/v1/import/docx/", {"file": file, "book": book.id, "mode": "split_headings", "split_level": 1}, format="multipart")
    assert [c["title"] for c in r.json()] == ["Частина перша"]
    file.seek(0)
    r = client.post("/api/v1/import/docx/preview/", {"file": file, "split_level": 3}, format="multipart")
    assert [c["title"] for c in r.json()["chapters"]] == ["rukopys"]


def test_docx_import_replace_chapter_and_errors(client, book, chapter, other_client, settings):
    file = make_docx([(None, "Новий текст розділу.", False)])
    r = client.post("/api/v1/import/docx/", {"file": file, "book": book.id, "mode": "replace_chapter", "chapter": chapter.id}, format="multipart")
    assert r.status_code == 201
    chapter.refresh_from_db()
    assert chapter.plain_text == "Новий текст розділу." and chapter.versions.filter(reason="import").exists()

    broken = io.BytesIO(b"PK\x03\x04 not really a zip")
    broken.name = "broken.docx"
    r = client.post("/api/v1/import/docx/", {"file": broken, "book": book.id}, format="multipart")
    assert r.status_code == 400 and r.json()["error"]["code"] == "validation_error"

    empty = make_docx([])
    r = client.post("/api/v1/import/docx/", {"file": empty, "book": book.id}, format="multipart")
    assert r.status_code == 400 and "тексту" in r.json()["error"]["message"]

    r = other_client.post("/api/v1/import/docx/", {"file": make_docx([(None, "x", False)]), "book": book.id}, format="multipart")
    assert r.status_code == 404

    settings.WW_MAX_DOCX_BYTES = 1000
    r = client.post("/api/v1/import/docx/", {"file": make_docx([(None, "x" * 5000, False)]), "book": book.id}, format="multipart")
    assert r.status_code == 400


def test_docx_export_roundtrip(client, book, chapter):
    chapter.content = {"type": "doc", "content": [
        {"type": "heading", "attrs": {"level": 2}, "content": [{"type": "text", "text": "Підзаголовок"}]},
        {"type": "paragraph", "content": [{"type": "text", "text": "Курсив", "marks": [{"type": "italic"}]}, {"type": "hardBreak"}, {"type": "text", "text": "далі"}]},
    ]}
    chapter.save()
    r = client.get(f"/api/v1/books/{book.id}/export/")
    assert r.status_code == 200
    assert r["Content-Type"].startswith("application/vnd.openxmlformats")
    exported = Document(io.BytesIO(r.content))
    texts = [p.text for p in exported.paragraphs]
    assert "Сяйво безодні" in texts and "Розділ 1" in texts and "Підзаголовок" in texts
    italic_runs = [run for p in exported.paragraphs for run in p.runs if run.italic]
    assert italic_runs and italic_runs[0].text.strip() == "Курсив"
    r = client.get(f"{chapters_url(book)}{chapter.id}/export/")
    assert r.status_code == 200 and "filename" in r["Content-Disposition"]


def test_manuscript_endpoint(client, book, chapter):
    data = client.get(f"/api/v1/books/{book.id}/manuscript/").json()
    assert data["chapters"][0]["text"].startswith("Туман")


def test_arcs_group_a_range_of_chapters(client, book, other_client, user):
    ids = [client.post(chapters_url(book), {"title": f"Розділ {i}"}, format="json").json()["id"] for i in range(1, 6)]
    arcs = f"/api/v1/books/{book.id}/arcs/"
    r = client.post(arcs, {"title": "Частина перша", "first_chapter": ids[3], "last_chapter": ids[1]}, format="json")
    assert r.status_code == 201
    first = r.json()["id"]
    by_id = {c["id"]: c["arc"] for c in client.get(chapters_url(book)).json()["results"]}
    assert [by_id[i] for i in ids] == [None, first, first, first, None]  # the range in reading order, either way round

    # a new range moves chapters out of the old one; another arc takes what it covers
    client.patch(f"{arcs}{first}/", {"first_chapter": ids[0], "last_chapter": ids[1]}, format="json")
    second = client.post(arcs, {"title": "Частина друга", "first_chapter": ids[1], "last_chapter": ids[4]}, format="json").json()["id"]
    by_id = {c["id"]: c["arc"] for c in client.get(chapters_url(book)).json()["results"]}
    assert [by_id[i] for i in ids] == [first, second, second, second, second]

    # a chapter can be put into an arc directly; deleting an arc keeps its chapters
    assert client.patch(f"{chapters_url(book)}{ids[0]}/", {"arc": second}, format="json").json()["arc"] == second
    assert client.delete(f"{arcs}{second}/").status_code == 204
    assert Chapter.objects.filter(book=book, arc__isnull=False).count() == 0 and Chapter.objects.filter(book=book).count() == 5

    # only both ends together, only this book's chapters, and others cannot see or use the arcs
    assert client.post(arcs, {"title": "x", "first_chapter": ids[0]}, format="json").status_code == 400
    other_book = Book.objects.create(world=book.world, title="Інша")
    alien = Chapter.objects.create(book=other_book, title="Чужий")
    assert client.post(arcs, {"title": "x", "first_chapter": alien.id, "last_chapter": ids[0]}, format="json").status_code == 400
    assert not book.arcs.filter(title="x").exists()
    foreign_arc = other_book.arcs.create(title="Чужа арка")
    assert client.patch(f"{chapters_url(book)}{ids[0]}/", {"arc": foreign_arc.id}, format="json").status_code == 400
    assert other_client.get(arcs).status_code == 404
